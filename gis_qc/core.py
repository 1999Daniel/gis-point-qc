"""仅读取输入，按明确配置核查点位属性和十进制度坐标。"""
from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import PurePosixPath

from defusedxml import ElementTree as SafeET
from openpyxl import load_workbook

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
HEADERS = ("编号", "名称", "类别", "经度", "纬度")
DEFAULT_CATEGORIES = ("学校", "医疗", "交通", "文体", "公共服务")
RULES = [
    {"code": code, "title": title} for code, title in (
        ("REQUIRED", "必填项"), ("DUPLICATE_ID", "编号唯一性"),
        ("TYPE_TEXT", "文本类型"), ("NUMBER", "有限数值"),
        ("RANGE", "经纬度范围"), ("CATEGORY", "类别白名单"),
        ("FORMULA", "禁止公式"), ("COORD_SWAP", "疑似经纬度颠倒"),
    )
]


class InputError(ValueError):
    """文件结构或配置无法安全处理。"""


class _ExcelError:
    """保留 Excel 错误单元格的类型，避免 #REF! 被当成正常文本。"""

    def __init__(self, value):
        self.value = value

    def __str__(self):
        return str(self.value)


def _text(value):
    return "" if value is None else str(value)


def _blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def _header(values):
    names = [_text(v).strip() for v in values]
    if len(set(names)) != len(names):
        raise InputError("表头存在重复字段，请使用且仅使用：" + "、".join(HEADERS))
    if len(names) != 5 or set(names) != set(HEADERS):
        raise InputError("表头缺失或存在额外字段；必须且仅包含：" + "、".join(HEADERS))
    return [names.index(name) for name in HEADERS]


def _read_csv(data):
    encoding = "UTF-8（支持 BOM）"
    try:
        content = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            content = data.decode("gb18030")
            encoding = "GB18030（UTF-8 解码失败后回退）"
        except UnicodeDecodeError as exc:
            raise InputError("CSV 编码无法识别，请另存为 UTF-8 CSV。") from exc
    if "\x00" in content:
        raise InputError("CSV 含空字节，可能不是文本 CSV。")
    reader = csv.reader(io.StringIO(content, newline=""), strict=True)
    rows, ignored = [], 0
    try:
        order = _header(next(reader))
        count = 0
        while True:
            row_number = reader.line_num + 1
            try:
                values = next(reader)
            except StopIteration:
                break
            if all(_blank(v) for v in values):
                ignored += 1
                continue
            count += 1
            if count > MAX_ROWS:
                raise InputError("数据行超过 10000 行限制。")
            if len(values) != 5:
                raise InputError(f"CSV 第 {row_number} 行有 {len(values)} 列，应为 5 列。")
            values = [values[index] for index in order]
            # CSV 不执行公式；以等号开头的公式文本同样作为问题标记。
            formulas = {i for i, value in enumerate(values) if value.lstrip().startswith("=")}
            rows.append((row_number, values, formulas))
    except StopIteration as exc:
        raise InputError("文件为空，缺少表头。") from exc
    except csv.Error as exc:
        raise InputError(f"CSV 格式错误（物理行 {reader.line_num}）：{exc}") from exc
    return rows, ignored, encoding


def _read_xlsx(data):
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    # 先约束展开体积和实际单元格；不信任 dimension 声明。
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or len({m.filename for m in members}) != len(members):
                raise InputError("XLSX 压缩包条目异常或存在重名条目。")
            if sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
                raise InputError("XLSX 解压后的体积超过 20 MiB 限制。")
            row_numbers = []
            for member in members:
                if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                    xml = SafeET.fromstring(archive.read(member))
                    actual_rows = xml.findall(f"{ns}sheetData/{ns}row")
                    if len(actual_rows) > 50000:
                        raise InputError("XLSX 存储行超过 50000 行，请清理无用空白格式行。")
                    seen_rows, seen_cells = set(), set()
                    numbers = []
                    for row in actual_rows:
                        rn = int(row.attrib.get("r", "0"))
                        if not 1 <= rn <= 1048576 or rn in seen_rows:
                            raise InputError("XLSX 行号异常或重复。")
                        seen_rows.add(rn)
                        numbers.append(rn)
                        for cell in row.findall(f"{ns}c"):
                            address = cell.attrib.get("r", "")
                            match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", address)
                            if not match or int(match[2]) != rn or address in seen_cells:
                                raise InputError("XLSX 单元格地址异常或重复。")
                            if match[1] not in ("A", "B", "C", "D", "E"):
                                raise InputError("XLSX 存在第 5 列以外的单元格，请清理额外列及其格式。")
                            seen_cells.add(address)
                    row_numbers.append(numbers)
            if len(row_numbers) != 1:
                raise InputError("XLSX 必须只包含一个工作表。")
        workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
        try:
            if len(workbook.sheetnames) != 1 or len(workbook.worksheets) != 1:
                raise InputError("XLSX 必须只包含一个工作表。")
            sheet = workbook.worksheets[0]
            if sheet.merged_cells.ranges:
                raise InputError("XLSX 不支持合并单元格，请取消合并后重试。")
            order = _header([sheet.cell(1, col).value for col in range(1, 6)])
            rows, ignored = [], 0
            for rn in sorted(row_numbers[0]):
                if rn == 1:
                    continue
                cells = [sheet.cell(rn, index + 1) for index in order]
                values = [_ExcelError(cell.value) if cell.data_type == "e" else cell.value for cell in cells]
                if all(_blank(value) for value in values):
                    ignored += 1
                    continue
                formulas = {i for i, cell in enumerate(cells) if cell.data_type == "f"}
                rows.append((rn, values, formulas))
                if len(rows) > MAX_ROWS:
                    raise InputError("数据行超过 10000 行限制。")
            return rows, ignored, "XLSX（读取原始单元格，公式不执行）"
        finally:
            workbook.close()
    except InputError:
        raise
    except Exception as exc:
        raise InputError("XLSX 文件损坏、格式不受支持或包含不安全的 XML。") from exc


def check_bytes(data: bytes, filename: str, crs="未声明", allowed_categories=None, is_demo=False) -> dict:
    """检查上传字节并返回可序列化结果；不改写文件，不进行坐标转换。"""
    if not isinstance(data, bytes) or not data:
        raise InputError("文件为空或内容不是字节数据。")
    if len(data) > MAX_BYTES:
        raise InputError("文件超过 5 MiB 大小限制。")
    if crs not in ("WGS84", "CGCS2000", "未声明"):
        raise InputError("坐标参考系只能选择 WGS84、CGCS2000 或未声明。")
    categories = DEFAULT_CATEGORIES if allowed_categories is None else allowed_categories
    if not isinstance(categories, (list, tuple)) or not 1 <= len(categories) <= 100:
        raise InputError("类别白名单须为 1 至 100 个文本值。")
    if any(not isinstance(v, str) or not v.strip() or len(v) > 100 or any(ord(c) < 32 for c in v) for v in categories):
        raise InputError("每个类别须为不含控制字符的非空文本，长度不超过 100。")
    categories = [v.strip() for v in categories]
    if len(set(categories)) != len(categories):
        raise InputError("类别白名单存在重复项（忽略首尾空白）。")
    filename = str(filename).replace("\\", "/").rsplit("/", 1)[-1]
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix == ".csv":
        rows, ignored, encoding = _read_csv(data)
    elif suffix == ".xlsx":
        rows, ignored, encoding = _read_xlsx(data)
    else:
        raise InputError("仅支持 .csv 和 .xlsx 文件。")
    issues = []
    ids = Counter(_text(v[0]).strip() for _, v, f in rows if not _blank(v[0]) and 0 not in f)
    for rn, values, formulas in rows:
        identifier = _text(values[0]).strip()
        def add(index, code, message, suggestion, severity="error"):
            issues.append({"row": rn, "id": identifier, "field": HEADERS[index] if index is not None else "经度/纬度",
                           "code": code, "severity": severity, "value": _text(values[index]) if index is not None else f"{_text(values[3])}, {_text(values[4])}",
                           "message": message, "suggestion": suggestion})
        for i, value in enumerate(values):
            if i in formulas:
                add(i, "FORMULA", "单元格含公式，未执行或使用缓存结果。", "核实原始数据后，替换为明确的文本或数值。")
            elif _blank(value):
                add(i, "REQUIRED", "必填字段为空。", "根据原始记录补充此字段。")
            elif i <= 2 and not isinstance(value, str):
                message = "此单元格为 Excel 错误值，不能作为正常文本。" if isinstance(value, _ExcelError) else "此字段必须是文本类型。"
                add(i, "TYPE_TEXT", message, "核实原始值后设为文本；编号丢失的前导零无法自动恢复。")
        if not _blank(values[0]) and 0 not in formulas and ids[identifier] > 1:
            add(0, "DUPLICATE_ID", "编号在文件内重复（忽略首尾空白）。", "核对全部同编号记录并提供唯一编号。")
        if not _blank(values[2]) and 2 not in formulas and isinstance(values[2], str) and values[2].strip() not in categories:
            add(2, "CATEGORY", "类别不在当前白名单中。", "核实类别或调整本次白名单。")
        numbers = {}
        for i, bound in ((3, 180), (4, 90)):
            value = values[i]
            if _blank(value) or i in formulas:
                continue
            try:
                if isinstance(value, bool):
                    raise ValueError
                number = float(value)
                if not math.isfinite(number):
                    raise ValueError
            except (ValueError, TypeError, OverflowError):
                add(i, "NUMBER", "坐标必须是有限数值（十进制度）。", "核实坐标格式，排除文字、布尔值、NaN 和无穷值。")
                continue
            numbers[i] = number
            if not -bound <= number <= bound:
                add(i, "RANGE", f"坐标超出 [-{bound}, {bound}] 度范围。", "核对原始坐标、单位及经纬度列。")
        if 3 in numbers and 4 in numbers:
            lon, lat = numbers[3], numbers[4]
            if not (-180 <= lon <= 180 and -90 <= lat <= 90) and -180 <= lat <= 180 and -90 <= lon <= 90:
                add(None, "COORD_SWAP", "交换经纬度后范围合法，疑似列顺序颠倒；尚未确认。", "人工对照数据源确认，系统不会自动交换。", "warning")
    counts = Counter(issue["code"] for issue in issues)
    errors = sum(issue["severity"] == "error" for issue in issues)
    notes = [
        "仅完成本次配置的字段、类型、唯一性、类别和坐标范围检查；真实位置、空间精度与数据来源未核验。",
        "坐标按十进制度解释；WGS84 / CGCS2000 为用户声明，不代表已鉴定，不进行坐标转换。",
        f"输入读取方式：{encoding}。编号不恢复前导零；文本比较忽略首尾空白。",
        f"已忽略 {ignored} 条完全空白记录；保留原始物理行号（CSV 多行字段取记录起始行）。XLSX 未存储的空行不计入忽略数。",
        "限制：文件不超过 5 MiB；非空数据不超过 10000 行；XLSX 解压不超过 20 MiB、存储行不超过 50000，单工作表、5 列、无合并单元格。",
    ]
    if crs == "未声明":
        notes.append("坐标参考系未声明，坐标参考系适用性待核实。")
    if not rows:
        notes.append("没有可检查的数据行，不能据此认定数据质量合格。")
    return {"schema_version": "1.0", "filename": filename, "sha256": hashlib.sha256(data).hexdigest(),
            "generated_at": datetime.now(timezone.utc).isoformat(), "crs": crs, "is_demo": bool(is_demo),
            "allowed_categories": categories,
            "summary": {"total_rows": len(rows), "ignored_blank_rows": ignored, "issue_count": len(issues),
                        "affected_rows": len({i["row"] for i in issues}), "error_count": errors,
                        "warning_count": len(issues) - errors, "by_rule": dict(counts),
                        "status": "发现问题" if issues else "待核实" if crs == "未声明" or not rows else "检查通过"},
            "issues": issues, "rules": [dict(rule) for rule in RULES], "notes": notes}
