"""重建人工模拟测试样例；不读取外部 POI，也不覆盖用户输入。"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ["编号", "名称", "类别", "经度", "纬度"]
CATEGORIES = ["学校", "医疗", "交通", "文体", "公共服务"]
NORMAL = [
    [f"{i:04d}", f"模拟点位{i:02d}", CATEGORIES[(i - 1) % 5], f"113.{100000 + i:06d}", f"23.{100000 + i:06d}"]
    for i in range(1, 13)
]
ISSUES = [
    ["0001", "", "学校", "113", "23"],
    ["0002", "模拟重复编号甲", "学校", "113", "23"],
    ["0002", "模拟重复编号乙", "医疗", "113", "23"],
    ["0004", "模拟非法类别", "商业", "113", "23"],
    ["0005", "模拟非数字经度", "交通", "东经113", "23"],
    ["0006", "模拟经度越界", "文体", "181", "23"],
    ["0007", "模拟纬度越界", "医疗", "113", "-91"],
    ["0008", "模拟疑似经纬度反置", "交通", "23", "113"],
    ["0009", "模拟缺失经度", "文体", "", "23"],
    ["", "模拟缺失编号", "公共服务", "113", "23"],
    ["0011", "模拟正常对照甲", "学校", "113", "23"],
    ["0012", "模拟正常对照乙", "公共服务", "113", "23"],
]


def build(*, template_only: bool = False) -> None:
    destination = ROOT / "samples"
    destination.mkdir(parents=True, exist_ok=True)
    datasets = [("template", [])]
    if not template_only:
        datasets.extend((("normal", NORMAL), ("issues", ISSUES)))
    for stem, rows in datasets:
        with (destination / f"{stem}.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(HEADERS)
            writer.writerows(rows)
        book = Workbook()
        sheet = book.active
        sheet.title = "模拟点位"
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        # 编号必须作为文本保存：保留 0001 等前导零。
        for cells in sheet.iter_rows(min_row=2):
            cells[0].number_format = "@"
        sheet.column_dimensions["A"].number_format = "@"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(color="FFFFFF", bold=True)
        for column, width in {"A": 15, "B": 28, "C": 16, "D": 18, "E": 18}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        book.save(destination / f"{stem}.xlsx")
        book.close()
        print(f"人工模拟样例 {stem}: CSV + XLSX，各 {len(rows)} 条")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="重建人工模拟样例与空白模板")
    parser.add_argument("--template-only", action="store_true", help="只生成空白模板，不改写现有模拟样例")
    build(template_only=parser.parse_args().template_only)
