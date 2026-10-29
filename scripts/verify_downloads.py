"""独立回读浏览器实际下载的成果。"""
import csv, hashlib, json
from pathlib import Path
from openpyxl import load_workbook
root = Path(__file__).resolve().parents[1]
downloads = root / "verification/downloads"
expected = json.loads((root / "samples/expected.json").read_text(encoding="utf-8"))["issues"]
raw = (downloads / "browser-issues.csv").read_bytes()
assert raw.startswith(b"\xef\xbb\xbf")
rows = list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))
actual = sorted((int(r["原始行号"]), r["字段"], r["规则代码"], r["严重程度"]) for r in rows)
assert actual == sorted(tuple(i) for i in expected["issues"])
html = (downloads / "browser-report.html").read_text(encoding="utf-8")
assert hashlib.sha256((root / "samples/issues.xlsx").read_bytes()).hexdigest() in html
assert "模拟样例" in html and "真实位置" in html
for suffix in ["csv", "xlsx"]:
    assert (downloads / f"template.{suffix}").read_bytes() == (root / "samples" / f"template.{suffix}").read_bytes()
book = load_workbook(downloads / "template.xlsx")
assert [book.active.cell(1,c).value for c in range(1,6)] == ["编号","名称","类别","经度","纬度"]
assert book.active.max_row == 1
book.close()
summary = {"passed": True,"issue_rows":len(rows),"exact_manual_expectations":True,"html_source_hash_matches":True,"downloaded_templates_match_originals":True}
(root / "verification/download-verification.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(summary,ensure_ascii=False))
