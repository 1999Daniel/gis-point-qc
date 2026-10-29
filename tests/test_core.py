"""使用人工预期、边界输入和恶意显示文本验证首版行为。"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

from gis_qc.core import InputError, check_bytes
from gis_qc.report import build_report, export_csv

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ["编号", "名称", "类别", "经度", "纬度"]
GOOD = ["0001", "模拟点位", "学校", "113", "23"]


def csv_bytes(rows, headers=HEADERS):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows, *, second_sheet=False, headers=HEADERS):
    book = Workbook()
    sheet = book.active
    sheet.title = "模拟点位"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if second_sheet:
        book.create_sheet("第二表")
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def tuples(result):
    return sorted((item["row"], item["field"], item["code"], item["severity"]) for item in result["issues"])


class ManualFixtureTests(unittest.TestCase):
    """预期答案来自手工清单，禁止用实现生成答案。"""

    @classmethod
    def setUpClass(cls):
        cls.expected = json.loads((ROOT / "samples" / "expected.json").read_text(encoding="utf-8"))

    def verify_fixture(self, stem, suffix):
        path = ROOT / "samples" / f"{stem}.{suffix}"
        data = path.read_bytes()
        before = hashlib.sha256(data).hexdigest()
        result = check_bytes(data, path.name, crs="WGS84", is_demo=True)
        expected = self.expected[stem]
        self.assertEqual(tuples(result), sorted(tuple(item) for item in expected["issues"]))
        for field in ("total_rows", "affected_rows", "issue_count", "error_count", "warning_count"):
            self.assertEqual(result["summary"][field], expected[field], field)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)
        self.assertEqual(hashlib.sha256(data).hexdigest(), before)
        self.assertEqual(result["summary"]["ignored_blank_rows"], 0)
        return result

    def test_normal_csv_matches_manual_expectations(self):
        self.verify_fixture("normal", "csv")

    def test_normal_xlsx_matches_manual_expectations(self):
        self.verify_fixture("normal", "xlsx")

    def test_issue_csv_matches_every_manual_issue(self):
        self.verify_fixture("issues", "csv")

    def test_issue_xlsx_matches_every_manual_issue(self):
        self.verify_fixture("issues", "xlsx")

    def test_fixture_storage_preserves_text_identifiers(self):
        with (ROOT / "samples" / "normal.csv").open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(next(csv.DictReader(stream))["编号"], "0001")
        book = load_workbook(ROOT / "samples" / "normal.xlsx", read_only=True)
        self.assertEqual(book.sheetnames, ["模拟点位"])
        self.assertEqual(book.active["A2"].value, "0001")
        self.assertEqual(book.active["A2"].data_type, "s")
        book.close()


class RuleBehaviorTests(unittest.TestCase):
    def test_undeclared_crs_stays_pending_with_no_field_issues(self):
        result = check_bytes(csv_bytes([GOOD]), "normal.csv")
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["summary"]["status"], "待核实")

    def test_declared_crs_with_no_errors_passes_table_checks(self):
        result = check_bytes(csv_bytes([GOOD]), "normal.csv", crs="WGS84")
        self.assertEqual(result["summary"]["status"], "检查通过")

    def test_duplicate_marks_all_members_and_keeps_zero_prefix(self):
        rows = [["0001", f"模拟{i}", "学校", "113", "23"] for i in range(3)]
        result = check_bytes(csv_bytes(rows), "duplicates.csv")
        self.assertEqual(tuples(result), [(row, "编号", "DUPLICATE_ID", "error") for row in (2, 3, 4)])
        self.assertEqual([issue["id"] for issue in result["issues"]], ["0001"] * 3)

    def test_text_identifiers_are_not_coerced_to_numbers(self):
        rows = [GOOD, ["1", "模拟另一点", "学校", "113", "23"]]
        self.assertEqual(check_bytes(csv_bytes(rows), "text_ids.csv")["issues"], [])

    def test_numeric_xlsx_id_is_rejected_even_if_display_format_has_zeros(self):
        book = Workbook()
        sheet = book.active
        sheet.append(HEADERS)
        sheet.append([1, *GOOD[1:]])
        sheet["A2"].number_format = "0000"
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        result = check_bytes(stream.getvalue(), "numeric_id.xlsx")
        self.assertEqual(tuples(result), [(2, "编号", "TYPE_TEXT", "error")])

    def test_missing_coordinates_only_report_required(self):
        result = check_bytes(csv_bytes([["0001", "模拟缺坐标", "学校", "", "   "]]), "empty_coord.csv")
        self.assertEqual(tuples(result), [(2, "纬度", "REQUIRED", "error"), (2, "经度", "REQUIRED", "error")])

    def test_each_required_field_can_be_detected(self):
        for index, field in enumerate(HEADERS):
            with self.subTest(field=field):
                row = GOOD.copy()
                row[index] = ""
                self.assertEqual(tuples(check_bytes(csv_bytes([row]), "missing.csv")), [(2, field, "REQUIRED", "error")])

    def test_blank_rows_are_ignored_without_reindexing_source_rows(self):
        result = check_bytes(csv_bytes([["", "", "", "", ""], ["0001", "", "学校", "113", "23"], [" "] * 5]), "blank.csv")
        self.assertEqual(result["summary"]["total_rows"], 1)
        self.assertEqual(result["summary"]["ignored_blank_rows"], 2)
        self.assertEqual(tuples(result), [(3, "名称", "REQUIRED", "error")])

    def test_nonfinite_coordinates_are_number_errors(self):
        for value in ("NaN", "nan", "Inf", "-Infinity", "1e999"):
            with self.subTest(value=value):
                row = ["0001", "模拟非有限数", "学校", value, "23"]
                self.assertEqual(tuples(check_bytes(csv_bytes([row]), "nonfinite.csv")), [(2, "经度", "NUMBER", "error")])

    def test_decimal_boundaries_are_inclusive(self):
        rows = [["0001", "模拟正边界", "学校", "180", "90"], ["0002", "模拟负边界", "学校", "-180", "-90"]]
        self.assertEqual(check_bytes(csv_bytes(rows), "boundaries.csv")["issues"], [])

    def test_out_of_range_without_valid_swap_is_not_swap_candidate(self):
        result = check_bytes(csv_bytes([["0001", "模拟双越界", "学校", "181", "91"]]), "outside.csv")
        self.assertEqual(tuples(result), [(2, "纬度", "RANGE", "error"), (2, "经度", "RANGE", "error")])

    def test_legal_coordinate_pair_does_not_trigger_swap_guess(self):
        result = check_bytes(csv_bytes([["0001", "模拟小经度", "学校", "23", "45"]]), "legal_pair.csv")
        self.assertEqual(result["issues"], [])

    def test_custom_allowed_categories_are_used(self):
        row = ["0001", "模拟设施", "雨水井", "113", "23"]
        result = check_bytes(csv_bytes([row]), "custom.csv", allowed_categories=["雨水井", "路灯"])
        self.assertEqual(result["issues"], [])
        self.assertEqual(tuples(check_bytes(csv_bytes([GOOD]), "custom.csv", allowed_categories=["雨水井"])), [(2, "类别", "CATEGORY", "error")])

    def test_formula_cells_have_no_secondary_type_category_or_number_issue(self):
        for index, field in enumerate(HEADERS):
            with self.subTest(field=field):
                row = GOOD.copy()
                row[index] = "=1+1"
                result = check_bytes(xlsx_bytes([row]), "formula.xlsx")
                self.assertEqual(tuples(result), [(2, field, "FORMULA", "error")])

    def test_excel_error_cells_never_pass_as_normal_text(self):
        for index, field in enumerate(HEADERS):
            with self.subTest(field=field):
                row = GOOD.copy()
                row[index] = "#REF!"
                result = check_bytes(xlsx_bytes([row]), "error_cell.xlsx", crs="WGS84")
                expected_code = "TYPE_TEXT" if index < 3 else "NUMBER"
                self.assertEqual(tuples(result), [(2, field, expected_code, "error")])
                self.assertEqual(result["summary"]["status"], "发现问题")

    def test_csv_formula_text_is_reported_without_execution(self):
        row = ["0001", "模拟公式", "学校", "=100+13", "23"]
        self.assertEqual(tuples(check_bytes(csv_bytes([row]), "formula.csv")), [(2, "经度", "FORMULA", "error")])

    def test_issue_records_include_actionable_context(self):
        result = check_bytes(csv_bytes([["0001", "模拟问题", "商业", "113", "23"]]), "context.csv")
        issue = result["issues"][0]
        self.assertEqual(issue["value"], "商业")
        self.assertTrue({"row", "id", "field", "code", "severity", "value", "message", "suggestion"}.issubset(issue))
        self.assertTrue(issue["message"])
        self.assertTrue(issue["suggestion"])


class InputRejectionTests(unittest.TestCase):
    def test_missing_header_is_rejected(self):
        with self.assertRaises(InputError):
            check_bytes(csv_bytes([["0001", "模拟点"]], headers=["编号", "名称"]), "missing_header.csv")

    def test_duplicate_header_is_rejected(self):
        with self.assertRaises(InputError):
            check_bytes(csv_bytes([[*GOOD, "extra"]], headers=[*HEADERS, "编号"]), "duplicate_header.csv")

    def test_multisheet_workbook_is_rejected(self):
        with self.assertRaises(InputError):
            check_bytes(xlsx_bytes([GOOD], second_sheet=True), "multi.xlsx")

    def test_damaged_xlsx_is_rejected(self):
        with self.assertRaises(InputError):
            check_bytes(b"this is not an xlsx zip", "damaged.xlsx")

    def test_unclosed_csv_quote_is_rejected(self):
        data = ','.join(HEADERS) + '\n0001,"unclosed,学校,113,23\n'
        with self.assertRaises(InputError):
            check_bytes(data.encode("utf-8"), "malformed.csv")

    def test_legacy_xls_is_rejected(self):
        with self.assertRaises(InputError):
            check_bytes(b"placeholder", "legacy.xls")

    def test_gb18030_csv_is_supported_with_encoding_noted(self):
        result = check_bytes("编号,名称,类别,经度,纬度\n0001,模拟,学校,113,23".encode("gb18030"), "gbk.csv")
        self.assertEqual(result["issues"], [])
        self.assertTrue(any("GB18030" in note for note in result["notes"]))

    def test_damaged_encoding_is_rejected_with_input_error(self):
        with self.assertRaises(InputError):
            check_bytes(b'\xff\xff\xff', "damaged.csv")

    def test_empty_template_is_pending_not_passed(self):
        for suffix in ["csv", "xlsx"]:
            data = (ROOT / "samples" / f"template.{suffix}").read_bytes()
            result = check_bytes(data, f"template.{suffix}", crs="WGS84")
            self.assertEqual(result["summary"]["total_rows"], 0)
            self.assertEqual(result["summary"]["status"], "待核实")

    def test_xlsx_blank_format_rows_do_not_count_as_data(self):
        book = Workbook()
        sheet = book.active
        sheet.append(HEADERS)
        sheet.append(GOOD)
        for rn in range(3, 10004):
            sheet.cell(rn, 1).number_format = "@"
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        result = check_bytes(stream.getvalue(), "blank_formats.xlsx", crs="WGS84")
        self.assertEqual(result["summary"]["total_rows"], 1)
        self.assertEqual(result["summary"]["ignored_blank_rows"], 10001)
        self.assertEqual(result["issues"], [])

    def test_row_limit_rejects_csv_and_excel(self):
        rows = [[str(i), "模拟", "学校", "113", "23"] for i in range(10001)]
        for data, name in [(csv_bytes(rows), "many.csv"), (xlsx_bytes(rows), "many.xlsx")]:
            with self.subTest(filename=name), self.assertRaises(InputError):
                check_bytes(data, name, crs="WGS84")


class ReportOutputTests(unittest.TestCase):
    def test_html_escapes_input_and_filename(self):
        payload = '<script>alert("模拟")</script>'
        filename = '<img src=x onerror=alert(1)>.csv'
        result = check_bytes(csv_bytes([["0001", "模拟名称", payload, "113", "23"]]), filename)
        html = build_report(result)
        self.assertIsInstance(html, str)
        self.assertNotIn(payload, html)
        self.assertNotIn(filename, html)
        self.assertIn("&lt;script&gt;", html)

    def test_csv_has_bom_and_one_row_per_issue(self):
        result = check_bytes((ROOT / "samples" / "issues.csv").read_bytes(), "issues.csv")
        exported = export_csv(result)
        self.assertTrue(exported.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(exported.decode("utf-8-sig"))))
        self.assertEqual(len(rows) - 1, 11)

    def test_csv_neutralizes_formula_prefixes_in_exported_values(self):
        for value in ("=HYPERLINK(\"https://example.invalid\",\"模拟\")", "+cmd", "@SUM(1,1)", "-cmd", "\t=1+1"):
            with self.subTest(value=value):
                result = check_bytes(csv_bytes([["0001", "模拟名称", value, "113", "23"]]), "export.csv")
                parsed = list(csv.reader(io.StringIO(export_csv(result).decode("utf-8-sig"))))
                self.assertGreater(len(parsed), 1)
                for row in parsed:
                    for cell in row:
                        self.assertFalse(cell.lstrip().startswith(("=", "+", "-", "@")), cell)

    def test_report_retains_pending_crs_warning(self):
        result = check_bytes(csv_bytes([GOOD]), "pending.csv")
        html = build_report(result)
        self.assertIn("待核实", html)
        self.assertIn("未声明", html)


if __name__ == "__main__":
    unittest.main()
